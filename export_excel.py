# -*- coding: utf-8 -*-
"""
将截图目录中的朋友圈截图逐张 OCR，按段落合并、跨截图全局去重，导出 Excel。

用法：
  python export_excel.py --name 好友名

输出：output/<好友名>朋友圈文本.xlsx
  Sheet1 朋友圈文本：序号 / 时间标记 / 首次出现截图 / 内容
  Sheet2 OCR明细  ：截图 / 坐标x / 坐标y / 识别文本（供核对）

注意：TOP_RESERVED / BODY_X / PARA_GAP 是按 1080x2400 手机截图调的版面参数，
换其他分辨率的设备需要重新标定。
"""
import os
import re
from rapidocr_onnxruntime import RapidOCR
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, 'output')

# 运行时由 main 设置
CAPS_DIR = ''
XLSX = ''
NICKNAME = ''

TOP_RESERVED = 240     # 状态栏/导航/封面以下才是内容
BODY_X = 300           # 正文区左边界（时间列在 x<240 的左侧）
PARA_GAP = 135         # 正文行 y 间距超过此值视为两条朋友圈

NOISE_EXACT = {
    '赞', '评论', '回复', '删除', '复制', '转发', '收藏', '翻译', '提到了我',
    '全文', '收起', '发送给朋友', '保存图片', '撤回', '更多',
    '提到了', '视频号', '取消', '置顶',
}
PHONE = re.compile(r'1\d{10}')
VIDEO_CARD = re.compile(r'^.{1,20}的视频$')
LIKE_LINE = re.compile(r'^[\u4e00-\u9fa5A-Za-z·、,，\s]{2,40}(赞|等人)$')
TIME_RE = re.compile(
    r'^(今天|昨天|前天|大前天|星期.?|\d+小时前|\d+分钟前|\d+天前|\d+天'
    r'|\d{1,2}月\d{1,2}日.*|\d{4}年\d{1,2}月\d{1,2}日.*|\d{4}年\d{1,2}月.*|\d{1,2}:\d{2})$')
YEAR_ONLY = re.compile(r'^\d{4}年$')   # 吸顶年份标题，不作为帖子时间


def norm(t):
    return re.sub(r'\s+', '', t)


def is_noise(t):
    t = t.strip()
    if not t or t in NOISE_EXACT:
        return True
    # 昵称行：精确匹配，或去掉符号/空白后相等（OCR 常漏掉昵称中的 > < 等符号）
    if t == NICKNAME:
        return True
    nk_core = re.sub(r'[^\u4e00-\u9fa5A-Za-z0-9]', '', NICKNAME)
    t_core = re.sub(r'[^\u4e00-\u9fa5A-Za-z0-9]', '', t)
    if nk_core and t_core == nk_core and len(t) <= len(NICKNAME) + 2:
        return True
    if YEAR_ONLY.match(t):
        return True
    if PHONE.search(t) or VIDEO_CARD.match(t) or LIKE_LINE.match(t):
        return True
    if re.fullmatch(r'[\W_]+', t):
        return True
    return False


def is_time_mark(t):
    return bool(TIME_RE.match(t.strip()))


def main():
    global CAPS_DIR, NICKNAME, XLSX
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--name', required=True, help='好友微信昵称（用于过滤昵称行和文件命名）')
    parser.add_argument('--caps-dir', default=None, help='截图目录（默认 caps_<昵称>）')
    parser.add_argument('--out', default=None, help='输出 xlsx 路径')
    a = parser.parse_args()

    def safe(s):
        return re.sub(r'[\\/:*?"<>|]', '', s).strip() or 'friend'

    NICKNAME = a.name
    sn = safe(a.name)
    CAPS_DIR = a.caps_dir or os.path.join(BASE_DIR, 'caps_' + sn)
    XLSX = a.out or os.path.join(OUTPUT_DIR, f'{sn}朋友圈文本.xlsx')

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    if not os.path.isdir(CAPS_DIR):
        print(f'截图目录不存在：{CAPS_DIR}')
        return
    files = sorted(f for f in os.listdir(CAPS_DIR) if f.endswith('.png'))
    if os.environ.get('CAP_LIMIT'):
        files = files[:int(os.environ['CAP_LIMIT'])]
    ocr = RapidOCR()

    detail_rows = []          # 明细 sheet
    paragraphs = []           # [(norm, text, time, cap)]
    index = {}                # norm -> paragraphs 下标，便于包含替换

    for n, fn in enumerate(files, 1):
        result, _ = ocr(os.path.join(CAPS_DIR, fn))
        items = []
        if result:
            for box, text, score in result:
                if score < 0.55:
                    continue
                t = text.strip()
                if not t:
                    continue
                x = (box[0][0] + box[1][0]) / 2
                y = (box[0][1] + box[2][1]) / 2
                items.append((y, x, t))
        items.sort()
        for y, x, t in items:
            if y > TOP_RESERVED:
                detail_rows.append((fn, round(x), round(y), t))

        # 时间列（左侧）标记，按 y 排序
        marks = [(y, t) for y, x, t in items
                 if x < 240 and y > TOP_RESERVED and is_time_mark(t)]

        # 正文行：右侧区域，过滤噪声
        body = [(y, t) for y, x, t in items
                if x > BODY_X and y > TOP_RESERVED and not is_noise(t)]

        # 按 y 间距切段
        segs = []
        cur = []
        last_y = None
        for y, t in body:
            if last_y is not None and y - last_y > PARA_GAP and cur:
                segs.append(cur)
                cur = []
            cur.append((y, t))
            last_y = y
        if cur:
            segs.append(cur)

        for seg in segs:
            y0 = seg[0][0]
            text = ''.join(t for _, t in seg)
            # 关联时间：取 y <= 段首行 的最近一个左侧时间标记
            tm = ''
            for my, mt in marks:
                if my <= y0 + 40:
                    tm = mt
                else:
                    break
            key = norm(text)
            if len(key) < 2:
                continue

            # 全局去重：完全相同 / 互相包含则保留更长的
            hit = index.get(key)
            if hit is not None:
                continue
            absorbed = False
            for i, (k2, txt2, tm2, cap2) in enumerate(paragraphs):
                if key in k2:                  # 本段是旧段的子串
                    absorbed = True
                    break
                if k2 in key:                 # 旧段是本段子串 -> 替换为更长文本
                    paragraphs[i] = (key, text, tm or tm2, fn)
                    index.pop(k2, None)
                    index[key] = i
                    absorbed = True
                    break
            if not absorbed:
                index[key] = len(paragraphs)
                paragraphs.append((key, text, tm, fn))

        print(f'[{n}/{len(files)}] {fn} -> 累计帖子 {len(paragraphs)}')

    # ---------- 写 Excel ----------
    wb = Workbook()
    ws = wb.active
    ws.title = '朋友圈文本'
    header_fill = PatternFill('solid', fgColor='4472C4')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    headers = ['序号', '时间标记', '首次出现截图', '内容']
    ws.append(headers)
    for c in range(1, 5):
        cell = ws.cell(1, c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for i, (_, text, tm, fn) in enumerate(paragraphs, 1):
        ws.append([i, tm, fn, text])
    widths = [6, 14, 16, 90]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = 'A2'
    for row in ws.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical='top')

    ws2 = wb.create_sheet('OCR明细')
    ws2.append(['截图', 'x', 'y', '识别文本'])
    for c in range(1, 5):
        cell = ws2.cell(1, c)
        cell.fill = header_fill
        cell.font = header_font
    for r in detail_rows:
        ws2.append(list(r))
    for c, w in enumerate([16, 8, 8, 80], 1):
        ws2.column_dimensions[get_column_letter(c)].width = w
    ws2.freeze_panes = 'A2'

    wb.save(XLSX)

    # 同步导出纯文本，作为 wordCloud.py 的输入
    txt_path = os.path.join(OUTPUT_DIR, f'{sn}已处理.txt')
    with open(txt_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(text for _, text, _, _ in paragraphs))

    print(f'\n完成：去重后帖子 {len(paragraphs)} 条，OCR明细 {len(detail_rows)} 行')
    print(f'输出：{XLSX}')
    print(f'文本：{txt_path}')


if __name__ == '__main__':
    main()
